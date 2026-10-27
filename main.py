"""INJUL - Consolidado Maestro de CCTV.

Flask + PostgreSQL (estado actual + historial trazable) + Redis (caché opcional).
Cada cambio de estado u observación se guarda al instante y queda en el historial.
"""

import json
import logging
import os
from contextlib import contextmanager
from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape

import openpyxl
import psycopg2
import redis
from flask import (
    Flask,
    jsonify,
    redirect,
    render_template_string,
    request,
    send_file,
    url_for,
)
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.graphics.charts.barcharts import HorizontalBarChart
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

try:
    from zoneinfo import ZoneInfo

    ZONA_HORARIA = ZoneInfo("America/Santo_Domingo")
except Exception:  # noqa: BLE001 - si falta tzdata se usa UTC-4 fijo (RD no tiene horario de verano)
    from datetime import timedelta, timezone

    ZONA_HORARIA = timezone(timedelta(hours=-4), "AST")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("injul-cctv")

app = Flask(__name__)

DATABASE_URL = os.environ.get("DATABASE_URL")
REDIS_URL = os.environ.get("REDIS_URL")
CACHE_KEY = "cctv:datos_reporte"
CACHE_TTL = 300  # segundos (la caché también se invalida en cada escritura)

# --- REDIS (CACHÉ OPCIONAL: si falla, la app sigue funcionando sin caché) ---
cache = None
if REDIS_URL:
    try:
        cache = redis.Redis.from_url(
            REDIS_URL,
            socket_connect_timeout=2,
            socket_timeout=2,
            decode_responses=True,
        )
    except Exception:  # noqa: BLE001
        logger.exception("No se pudo configurar Redis; se continúa sin caché")
        cache = None

# --- CATÁLOGO OFICIAL DE CÁMARAS POR SUCURSAL ---
CATALOGO_CAMARAS = {
    "Caribbean Agregados": [
        "Frente Oficina", "Puerta desde CR", "Lado Atras", "IP Camera 04",
        "Frente Lado desde Caribbean", "Lado Atras esquina puerta", "Puerta lado norte",
        "Frente hacia lado norte", "Oficina Gabinete data", "IP Camera 10", "IP Camera 11",
        "IP Camera 12", "IP Camera 13", "IP Camera 14", "IP Camera 15", "IP Camera 16",
    ],
    "Marmotech Blue Mall": [
        "Show 2da planta", "Area Clientes 1ra planta", "Area 1ra planta puerta trasera",
        "2da Planta banos", "Area 1ra planta puerta trasera (2)", "Gabinete TI",
        "IP Camera 07", "2da Planta banos (2)", "IP Camera 09", "IP Camera 10",
        "IP Camera 11", "IP Camera 12", "IP Camera 13", "IP Camera 14", "IP Camera 15",
        "IP Camera 16",
    ],
    "Marmotech DCM": [
        "Puerta Principal", "Cocina", "IP Camera 03", "S1 Entrada Vehiculos",
        "S2 Pasillo lateral", "IP Camera 06", "1ra planta asensor visita", "IP Camera 09",
        "IP Camera 11", "segundo nivel banera", "seg piso griferia", "IP Camera 14",
        "IP Camera 15", "Entrada edificio", "S2 Salida", "IP Camera 18", "IP Camera 19",
        "5to piso onix miel", "S2 Salida V", "S2 Salida Planta E", "Entrada Calle Parqueos",
        "Frente Edificio", "Salida S1 Vehiculos", "S1 Entrada al S2",
        "S1 Entrada vista de frente", "S2 Area Piezas y Vehiculos", "IP Camera 30",
        "almacen griferia piiso 1", "piso 1 entrada plancha", "IP Camera 33",
        "IP Camera 34", "IP Camera 35", "IP Camera 36", "IP Camera 37", "IP Camera 38",
        "IP Camera 39", "IP Camera 40", "IP Camera 41", "IP Camera 42", "IP Camera 43",
        "IP Camera 45", "Oficina Almacen", "Entrada a cobro lobby", "IP Camera 49",
        "Salida Asensor c", "IP Camera 51", "IP Camera 52", "IP Camera 53", "IP Camera 54",
        "IP Camera 55", "IP Camera 56",
    ],
    "Marmotech Outlet RP": [
        "IP Camera 01", "IP Camera 02", "IP Camera 03", "IP Camera 04", "IP Camera 05",
        "IP Camera 06", "IP Camera 07", "IP Camera 08", "IP Camera 09", "IP Camera 10",
        "IP Camera 11", "IP Camera 12", "IP Camera 13", "IP Camera 14", "IP Camera 15",
        "IP Camera 16",
    ],
    "Marmotech Punta Cana": [
        "Puerta Entrada Principal", "Area trasera 1", "Area Trasera 2", "Area Trasera 5",
        "Area trasera 3", "Area Trasera 4", "Oficina 1", "Area Gabinete datos", "Almacen",
        "Area Lateral Derecha", "Zona Carga / Descarga", "Almacen 1", "Frente Edificacion",
        "Instalacion", "Entrada Oficinas", "Almacen 2 Escalera Subir Techo Oficina",
        "Almacen 3",
    ],
    "Sucursal San Cristóbal": [
        "Almacen Repuesto", "Parqueo 1 Tienda", "Entrada Tienda S", "Parqueo 2 Tienda",
        "Almacen 1 Tienda", "C/Entrada Tienda", "Almacen 2 Tienda", "Almacen 3 Tienda",
        "Calle Entrada 2", "Oficina Tienda S", "Parqueo 3 Tiend", "Almacen 4 Tienda",
        "Almacen 5 Tienda", "Area Caja Oficin", "MODULARES 1",
        "Entrada Mantenimiento norte", "MODULARES Area F", "MODULARES 4", "Alm. Gene 04",
        "Alm Gene 06", "Cuarto de Compresores", "Alm Gene 01", "Alm Gene 02", "MOTORES 2",
        "ENTRADA P1", "Puerta Entrada", "ARTECH 02", "PRODUCCION 1", "ARTECH 01",
        "Alm Gene 03", "PULIDORAS", "ARTECH03", "PULIDORA 02", "DATA CENTER 1",
        "Carga Area 1 Alm", "Puerta Entrada A", "Area Motores Alm", "Carga 2 Alm Gene",
        "Gasoil entrada /", "Area Parqueo 1 F", "Salida Fabrica A", "Gasoil Dispensad",
        "Gasoil", "Parqueo Fabrica", "Puerta Entrada G", "Garita Guardian",
        "Escritorio Garit", "Frente Puerta no", "Almacen Repuesto (2)", "Ing y Calidad",
        "Alm Gene 07", "Area de Planta", "ARTECH 1", "ARTECH 2", "Alm Gene 05",
        "IP Camera 57", "IP Camera 58", "IP Camera 59",
    ],
    "Marmotech Santiago 1": [
        "Almacen Repuesto", "Parqueo 1 Tienda", "Entrada Tienda 5", "Parqueo 2 Tienda",
        "Almacen 1 Tienda", "C/Entrada Tienda", "Almacen 2 Tienda", "Almacen 3 Tienda",
        "Calle Entrada 2", "Oficina Tienda S", "Parqueo 3 Tiend", "Almacen 4 Tienda",
        "Almacen 5 Tienda", "Area Caja Oficin", "MODULARES 1",
        "Entrada Mantenimiento norte",
    ],
    "Marmotech Santiago 2": [
        "IP Camera 01", "IP Camera 02", "IP Camera 03", "IP Camera 04", "IP Camera 05",
        "IP Camera 06", "IP Camera 07", "IP Camera 08", "IP Camera 09", "IP Camera 10",
        "IP Camera 11", "IP Camera 12", "IP Camera 13", "IP Camera 14", "IP Camera 15",
        "IP Camera 16",
    ],
    "Santiago 2 - Sección Muebles": [
        "MUEBLES", "ENTRADA", "COCINA", "GRIFERIA", "PARQUEO O-E", "BANOS", "PISOS",
        "SALA REUNION", "PARQUEO E-O", "COUNTER", "MASETEROS", "IP Camera 12",
        "IP Camera 13", "IP Camera 14", "IP Camera 15", "IP Camera 16",
    ],
}
# --- CATÁLOGO OFICIAL DEL MONITOREO CÍCLICO (formato institucional INJUL) ---
# Es la cuota operativa que un operador audita por turno (5 a 10 cámaras clave por
# sede), NO el chequeo masivo de todas las cámaras (eso es CATALOGO_CAMARAS).
# El orden y la cantidad de cuadros de cada sede replican el PDF oficial: las sedes
# grandes (San Cristóbal y DCM) llevan 10 cuadros (ciclo 1 = arriba, ciclo 2 = abajo);
# el resto lleva 5. No agregar cámaras de más: rompe el formato.
CATALOGO_CICLICO = {
    "Sucursal San Cristóbal": [
        "Entrada puerta 1 - entrada p1", "Área de diseño artech - artech02",
        "Área de diseño artech 2 - artech01", "Pulidoras", "Pulidora 02",
        "Carga almacén general - carga área 1 alm", "Carga almacén general 2",
        "Entrada a fábrica puerta 3 - salida fábrica A",
        "Garita punchado - puerta entrada G", "Calle principal puerta 2 - frente puerta",
    ],
    "Marmotech DCM": [
        "IP Camera 30", "Frente de la puerta, utensilios - IP Camera 32",
        "Parqueo de motores - IP Camera 34", "Terraza lado céntrico - IP Camera 36",
        "IP Camera 38", "Parqueo sótano 1 - IP Camera 40",
        "Pasillos de exhibidores parte céntrica - IP Camera 42", "IP Camera 44",
        "IP Camera 46", "IP Camera 48",
    ],
    "Marmotech Santiago 1": [
        "Calle frente a la empresa - IP Camera 02", "Parte trasera - IP Camera 04",
        "Calle principal - IP Camera 06", "Mostrador - IP Camera 08",
        "Almacén interior 2 - IP Camera 10",
    ],
    "Marmotech Santiago 2": [
        "Mostrador - Muebles", "Cocina", "Parqueo principal - Parqueo 0-E", "Pisos",
        "Parqueo E-O",
    ],
    "Marmotech Blue Mall": [
        "Show 2da Planta", "Área 1ra Planta Puerta Trasera",
        "Área 1ra Planta Puerta Trasera Norte", "Área Clientes 1ra Planta", "Gabinete TI",
    ],
    "Marmotech Outlet RP": [
        "Parqueo lado derecho - IP Camera 10", "Parte atrás, compresores A/C - IP Camera 12",
        "Almacén General - IP Camera 14", "Exhibidores de venta - IP Camera 16",
        "Cámara escalera lado izquierdo - IP Camera 02",
    ],
    "Marmotech Punta Cana": [
        "Área Trasera 1", "Área Trasera 5", "Área Trasera 4", "Área Gabinete Datos",
        "Puerta Entrada Principal",
    ],
    "Marmotech Caribbean": [
        "CB 6-7-8 Planta Eléctrica", "Taller Soldar", "Oficina Almacén Repuesto",
        "Almacén Repuesto", "Taller Fabricación",
    ],
    "Caribbean Agregados": [
        "Frente Lado desde Caribbean", "Lado Atrás esquina puerta", "Puerta lado norte",
        "Frente hacia lado norte", "Oficina Gabinete data",
    ],
}

ESTADOS = [
    ("Operativa", "Operativa (Sí)"),
    ("No visible", "No visible (Rojo)"),
    ("Por verificar", "Por verificar"),
]
ESTADOS_VALIDOS = {valor for valor, _ in ESTADOS}

MENSAJES = {
    "ticket_ok": ("ok", "Ticket registrado correctamente."),
    "sin_db": ("error", "La base de datos no está configurada (falta DATABASE_URL)."),
    "error": ("error", "No se pudo guardar. Revisa los registros en Railway."),
}

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>INJUL - Consolidado CCTV</title>
    <style>
        body { font-family: Arial, sans-serif; background-color: #f4f6f9; color: #333; margin: 0; padding: 15px; }
        .container { max-width: 950px; margin: 0 auto; background: #fff; padding: 20px; border-radius: 8px; box-shadow: 0 4px 10px rgba(0,0,0,0.1); }
        h1 { color: #1F497D; font-size: 20px; border-bottom: 2px solid #1F497D; padding-bottom: 8px; }
        h3 { color: #1F497D; margin-top: 5px; }
        .card { background: #f8f9fa; padding: 15px; margin-bottom: 15px; border-radius: 6px; border-left: 4px solid #1F497D; }
        .btn { display: inline-block; background: #1F497D; color: white; padding: 10px 15px; text-decoration: none; border-radius: 4px; font-weight: bold; margin-top: 5px; margin-right: 5px; border: none; cursor: pointer; font-size: 14px; }
        .btn-excel { background: #217346; }
        .btn-ticket { background: #D9534F; }
        .btn-gris { background: #6c757d; }
        .btn:hover { opacity: 0.9; }
        .form-group { margin-bottom: 10px; }
        label { display: block; font-weight: bold; font-size: 13px; margin-bottom: 3px; color: #444; }
        input, select, textarea { width: 100%; padding: 8px; box-sizing: border-box; border: 1px solid #ccc; border-radius: 4px; font-size: 14px; }
        table { width: 100%; border-collapse: collapse; margin-top: 15px; }
        th, td { border: 1px solid #ddd; padding: 8px; text-align: center; font-size: 12px; }
        th { background-color: #1F497D; color: white; }
        .estado-op td.celda-estado { background: #dff0d8; }
        .estado-nv td.celda-estado { background: #f2dede; }
        .estado-pv td.celda-estado { background: #fcf8e3; }
        .marca { font-size: 11px; color: #777; white-space: nowrap; min-width: 80px; }
        .marca.ok { color: #2d6a2d; }
        .marca.error { color: #a94442; font-weight: bold; }
        .aviso { padding: 10px 14px; border-radius: 4px; margin-bottom: 15px; font-size: 14px; }
        .aviso.ok { background: #dff0d8; color: #2d6a2d; }
        .aviso.error { background: #f2dede; color: #a94442; }
        .kpis { display: flex; flex-wrap: wrap; gap: 8px; margin-bottom: 15px; }
        .kpi { flex: 1; min-width: 110px; background: #f8f9fa; border-top: 3px solid #1F497D; padding: 8px; text-align: center; }
        .kpi b { display: block; font-size: 22px; color: #1F497D; }
        .kpi span { font-size: 11px; color: #666; }
        .tabla-scroll { overflow-x: auto; }
    </style>
    <script>
        var CLASES = {'Operativa': 'estado-op', 'No visible': 'estado-nv', 'Por verificar': 'estado-pv'};

        function mostrarSucursal() {
            var select = document.getElementById('sucursal-select');
            var listas = document.getElementsByClassName('lista-sucursal');
            for (var i = 0; i < listas.length; i++) {
                listas[i].style.display = (listas[i].id === 'sucursal-' + select.selectedIndex) ? 'block' : 'none';
            }
        }

        function actualizarResumen(r) {
            document.getElementById('k-total').textContent = r.total;
            document.getElementById('k-op').textContent = r.operativas;
            document.getElementById('k-nv').textContent = r.no_visibles;
            document.getElementById('k-pv').textContent = r.por_verificar;
            document.getElementById('k-disp').textContent = r.disponibilidad.toFixed(1) + '%';
        }

        function guardarFila(tr) {
            var marca = tr.querySelector('.marca');
            var estado = tr.querySelector('select').value;
            var obs = tr.querySelector('input').value;
            marca.className = 'marca';
            marca.textContent = 'Guardando...';
            fetch('/api/camara', {
                method: 'POST',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify({sucursal: tr.dataset.suc, idx: parseInt(tr.dataset.idx, 10), estado: estado, observacion: obs})
            }).then(function (resp) {
                return resp.json().then(function (j) { return {ok: resp.ok, j: j}; });
            }).then(function (res) {
                if (!res.ok || !res.j.ok) { throw new Error(res.j.error || 'Error'); }
                tr.className = 'camera-row ' + (CLASES[estado] || 'estado-pv');
                marca.className = 'marca ok';
                marca.textContent = res.j.cambiado ? 'Guardado \u2713' : 'Sin cambios';
                actualizarResumen(res.j.resumen);
            }).catch(function () {
                marca.className = 'marca error';
                marca.textContent = 'NO guardado: reintenta';
            });
        }

        window.onload = function () {
            mostrarSucursal();
            var filas = document.querySelectorAll('tr.camera-row');
            for (var i = 0; i < filas.length; i++) {
                (function (tr) {
                    tr.querySelector('select').addEventListener('change', function () { guardarFila(tr); });
                    tr.querySelector('input').addEventListener('change', function () { guardarFila(tr); });
                })(filas[i]);
            }
        };
    </script>
</head>
<body>
    <div class="container">
        <h1>INJUL - Sistema de Monitoreo CCTV</h1>
        <p style="margin: 0 0 12px;"><a class="btn btn-gris" href="/ciclico/">Plan de Monitoreo Cíclico y Novedades →</a></p>

        {% if mensaje %}<div class="aviso {{ tipo }}">{{ mensaje }}</div>{% endif %}

        <div class="kpis">
            <div class="kpi"><b id="k-total">{{ resumen.total }}</b><span>Cámaras</span></div>
            <div class="kpi"><b id="k-op">{{ resumen.operativas }}</b><span>Operativas</span></div>
            <div class="kpi"><b id="k-nv">{{ resumen.no_visibles }}</b><span>No visibles</span></div>
            <div class="kpi"><b id="k-pv">{{ resumen.por_verificar }}</b><span>Por verificar</span></div>
            <div class="kpi"><b id="k-disp">{{ '%.1f'|format(resumen.disponibilidad) }}%</b><span>Disponibilidad</span></div>
        </div>

        <div class="card">
            <h3>Seleccionar Sucursal y Revisar Cámaras</h3>
            <p style="font-size: 12px; color: #666; margin-top: 0;">
                Cada cambio de estado u observación se guarda al instante y queda registrado en el historial.
            </p>
            <div class="form-group">
                <label>Sucursal:</label>
                <select id="sucursal-select" onchange="mostrarSucursal()">
                    {% for suc in catalogo.keys() %}
                    <option value="{{ suc }}" {% if loop.index0 == seleccion %}selected{% endif %}>{{ suc }}</option>
                    {% endfor %}
                </select>
            </div>

            {% for suc, camaras in catalogo.items() %}
            {% set guardado = detalles.get(suc, {}) %}
            <div id="sucursal-{{ loop.index0 }}" class="lista-sucursal" style="display:none; margin-top: 15px;">
                <h4 style="color: #1F497D; margin-bottom: 5px;">Cámaras de {{ suc }}</h4>
                <div class="tabla-scroll">
                <table>
                    <tr>
                        <th>#</th>
                        <th>Nombre de la Cámara (EZStation)</th>
                        <th>Estado</th>
                        <th>Observación</th>
                        <th></th>
                    </tr>
                    {% for cam in camaras %}
                    {% set info = guardado.get(cam, {}) %}
                    {% set est = info.get('estado') or 'Por verificar' %}
                    <tr class="camera-row {{ clases[est] }}" data-suc="{{ suc }}" data-idx="{{ loop.index0 }}">
                        <td>{{ loop.index }}</td>
                        <td style="text-align: left; font-weight: bold;">{{ cam }}</td>
                        <td class="celda-estado">
                            <select style="padding: 4px;">
                                {% for valor, texto in estados %}
                                <option value="{{ valor }}" {% if est == valor %}selected{% endif %}>{{ texto }}</option>
                                {% endfor %}
                            </select>
                        </td>
                        <td>
                            <input type="text" value="{{ info.get('observacion', '') }}" maxlength="500" placeholder="Detalle si falla..." style="padding: 4px;">
                        </td>
                        <td class="marca"></td>
                    </tr>
                    {% endfor %}
                </table>
                </div>
            </div>
            {% endfor %}
        </div>

        <div class="card">
            <h3>Registrar Ticket o Novedad con Proveedor</h3>
            <form action="/guardar-ticket" method="POST">
                <div class="form-group">
                    <label>Número de Ticket:</label>
                    <input type="text" name="ticket" required maxlength="50" placeholder="Ej: 10437 o SIN TICKET">
                </div>
                <div class="form-group">
                    <label>Sucursal o Área Afectada:</label>
                    <input type="text" name="sucursal_area" required maxlength="150" placeholder="Ej: Marmotech DCM">
                </div>
                <div class="form-group">
                    <label>Estado del Ticket:</label>
                    <select name="estado_ticket">
                        <option value="En proceso">En proceso</option>
                        <option value="Terminado">Terminado</option>
                    </select>
                </div>
                <div class="form-group">
                    <label>Días Abiertos:</label>
                    <input type="number" name="dias_abiertos" value="0" min="0" placeholder="Ej: 5">
                </div>
                <div class="form-group">
                    <label>Detalle del Reporte (Información de lo que pasó):</label>
                    <textarea name="detalle" rows="2" placeholder="Ej: Problemas de visualización..."></textarea>
                </div>
                <button type="submit" class="btn btn-ticket">📝 Registrar Ticket</button>
            </form>
        </div>

        <div class="card">
            <h3>Generación de Reportes Oficiales</h3>
            <a href="/descargar-pdf" class="btn">📥 Descargar PDF Oficial</a>
            <a href="/descargar-excel" class="btn btn-excel">📊 Descargar Excel Matriz</a>
            <a href="/historial" class="btn btn-gris">🕘 Ver Historial de Cambios</a>
        </div>
    </div>
</body>
</html>
"""

HISTORIAL_TEMPLATE = """
<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>INJUL - Historial CCTV</title>
    <style>
        body { font-family: Arial, sans-serif; background: #f4f6f9; color: #333; margin: 0; padding: 15px; }
        .container { max-width: 1100px; margin: 0 auto; background: #fff; padding: 20px; border-radius: 8px; box-shadow: 0 4px 10px rgba(0,0,0,0.1); }
        h1 { color: #1F497D; font-size: 20px; border-bottom: 2px solid #1F497D; padding-bottom: 8px; }
        table { width: 100%; border-collapse: collapse; margin-top: 10px; }
        th, td { border: 1px solid #ddd; padding: 6px; font-size: 12px; text-align: left; }
        th { background: #1F497D; color: #fff; }
        select, .btn { padding: 8px; font-size: 14px; }
        .btn { background: #1F497D; color: #fff; text-decoration: none; border-radius: 4px; }
        .tabla-scroll { overflow-x: auto; }
    </style>
</head>
<body>
    <div class="container">
        <h1>Historial de cambios (últimos {{ filas|length }})</h1>
        <form method="GET" action="/historial">
            <select name="sucursal" onchange="this.form.submit()">
                <option value="">Todas las sucursales</option>
                {% for suc in catalogo.keys() %}
                <option value="{{ suc }}" {% if suc == filtro %}selected{% endif %}>{{ suc }}</option>
                {% endfor %}
            </select>
            <a class="btn" href="/">← Volver al panel</a>
        </form>
        <div class="tabla-scroll">
        <table>
            <tr><th>Fecha y hora</th><th>Sucursal</th><th>Cámara</th><th>Estado anterior</th><th>Estado nuevo</th><th>Observación anterior</th><th>Observación nueva</th><th>Origen</th></tr>
            {% for f in filas %}
            <tr>{% for campo in f %}<td>{{ campo }}</td>{% endfor %}</tr>
            {% else %}
            <tr><td colspan="8">Aún no hay cambios registrados.</td></tr>
            {% endfor %}
        </table>
        </div>
    </div>
</body>
</html>
"""


# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------
CLASES_ESTADO = {
    "Operativa": "estado-op",
    "No visible": "estado-nv",
    "Por verificar": "estado-pv",
}


def ahora():
    return datetime.now(ZONA_HORARIA) if ZONA_HORARIA else datetime.now()


def local_txt(momento):
    """Formatea un datetime de la base de datos en hora de Santo Domingo."""
    if momento is None:
        return ""
    if ZONA_HORARIA and momento.tzinfo:
        momento = momento.astimezone(ZONA_HORARIA)
    return momento.strftime("%d/%m/%Y %I:%M:%S %p")


def esc(texto):
    """Escapa texto para usarlo dentro de un Paragraph de ReportLab."""
    return escape(str(texto if texto is not None else ""))


def ip_cliente():
    reenviado = request.headers.get("X-Forwarded-For", "")
    return (reenviado.split(",")[0].strip() or request.remote_addr or "")[:60]


@contextmanager
def db_cursor():
    """Abre conexión a Postgres; confirma al salir bien y revierte si hay error."""
    conn = psycopg2.connect(DATABASE_URL, connect_timeout=5)
    try:
        with conn:
            with conn.cursor() as cur:
                yield cur
    finally:
        conn.close()


def inicializar_base_de_datos():
    if not DATABASE_URL:
        logger.warning("DATABASE_URL no está definida; no se guardarán datos")
        return
    try:
        with db_cursor() as cur:
            cur.execute("""
                CREATE TABLE IF NOT EXISTS reportes_cctv_detalles (
                    id SERIAL PRIMARY KEY,
                    fecha DATE NOT NULL,
                    sucursal VARCHAR(100) NOT NULL,
                    camara VARCHAR(150) NOT NULL,
                    estado VARCHAR(50) NOT NULL,
                    observacion TEXT
                );
                CREATE TABLE IF NOT EXISTS reportes_tickets (
                    id SERIAL PRIMARY KEY,
                    fecha DATE NOT NULL,
                    ticket VARCHAR(50) NOT NULL,
                    sucursal_area VARCHAR(150) NOT NULL,
                    estado VARCHAR(50) NOT NULL,
                    dias_abiertos INT DEFAULT 0,
                    detalle TEXT,
                    creado_en TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                );
                CREATE TABLE IF NOT EXISTS estado_camaras_actual (
                    sucursal VARCHAR(100) NOT NULL,
                    camara VARCHAR(150) NOT NULL,
                    estado VARCHAR(50) NOT NULL,
                    observacion TEXT NOT NULL DEFAULT '',
                    actualizado_en TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    PRIMARY KEY (sucursal, camara)
                );
                CREATE TABLE IF NOT EXISTS historial_cctv (
                    id BIGSERIAL PRIMARY KEY,
                    cambiado_en TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    sucursal VARCHAR(100) NOT NULL,
                    camara VARCHAR(150) NOT NULL,
                    estado_anterior VARCHAR(50),
                    estado_nuevo VARCHAR(50) NOT NULL,
                    observacion_anterior TEXT,
                    observacion_nueva TEXT,
                    origen VARCHAR(60)
                );
                CREATE INDEX IF NOT EXISTS idx_historial_fecha
                    ON historial_cctv (cambiado_en DESC);
                CREATE INDEX IF NOT EXISTS idx_historial_camara
                    ON historial_cctv (sucursal, camara);
            """)
            # Migración única: si el estado actual está vacío, parte del último
            # reporte diario guardado con la versión anterior de la app.
            cur.execute("""
                INSERT INTO estado_camaras_actual (sucursal, camara, estado, observacion)
                SELECT sucursal, camara, estado, COALESCE(observacion, '')
                FROM reportes_cctv_detalles
                WHERE fecha = (SELECT MAX(fecha) FROM reportes_cctv_detalles)
                  AND NOT EXISTS (SELECT 1 FROM estado_camaras_actual)
                ON CONFLICT (sucursal, camara) DO NOTHING
            """)
    except Exception:  # noqa: BLE001
        logger.exception("Error inicializando la base de datos")


# ----------------------------------------------------------------------------
# Caché (Redis) y carga de datos
# ----------------------------------------------------------------------------
def cache_leer():
    if cache is None:
        return None
    try:
        crudo = cache.get(CACHE_KEY)
        return json.loads(crudo) if crudo else None
    except Exception:  # noqa: BLE001
        logger.warning("Redis no disponible al leer; se consulta la base de datos")
        return None


def cache_guardar(datos):
    if cache is None:
        return
    try:
        cache.setex(CACHE_KEY, CACHE_TTL, json.dumps(datos))
    except Exception:  # noqa: BLE001
        logger.warning("Redis no disponible al guardar caché")


def cache_invalidar():
    if cache is None:
        return
    try:
        cache.delete(CACHE_KEY)
    except Exception:  # noqa: BLE001
        logger.warning("Redis no disponible al invalidar caché")


def cargar_datos():
    """Devuelve {'detalles': {sucursal: {camara: {...}}}, 'tickets': [[...]], 'actualizado': str}."""
    datos = cache_leer()
    if datos is not None:
        return datos

    datos = {"detalles": {}, "tickets": [], "actualizado": ""}
    if not DATABASE_URL:
        return datos
    try:
        with db_cursor() as cur:
            cur.execute("""
                SELECT sucursal, camara, estado, observacion, actualizado_en
                FROM estado_camaras_actual
            """)
            ultima = None
            for suc, cam, est, obs, momento in cur.fetchall():
                datos["detalles"].setdefault(suc, {})[cam] = {
                    "estado": est,
                    "observacion": obs or "",
                }
                if ultima is None or momento > ultima:
                    ultima = momento
            datos["actualizado"] = local_txt(ultima)
            cur.execute("""
                SELECT ticket, fecha, sucursal_area, estado, dias_abiertos, detalle
                FROM reportes_tickets
                ORDER BY id DESC
            """)
            datos["tickets"] = [
                [str(t[0]), str(t[1]), str(t[2]), str(t[3]), str(t[4]), str(t[5] or "")]
                for t in cur.fetchall()
            ]
        cache_guardar(datos)
    except Exception:  # noqa: BLE001
        logger.exception("Error leyendo datos de la base de datos")
    return datos


def cargar_historial(limite=200, sucursal=None):
    """Últimos cambios registrados, más recientes primero (siempre desde la base de datos)."""
    if not DATABASE_URL:
        return []
    consulta = """
        SELECT cambiado_en, sucursal, camara, estado_anterior, estado_nuevo,
               observacion_anterior, observacion_nueva, origen
        FROM historial_cctv
    """
    params = []
    if sucursal:
        consulta += " WHERE sucursal = %s"
        params.append(sucursal)
    consulta += " ORDER BY id DESC LIMIT %s"
    params.append(limite)
    try:
        with db_cursor() as cur:
            cur.execute(consulta, params)
            return [
                [local_txt(f[0]), f[1], f[2], f[3] or "(sin registro)", f[4], f[5] or "", f[6] or "", f[7] or ""]
                for f in cur.fetchall()
            ]
    except Exception:  # noqa: BLE001
        logger.exception("Error leyendo el historial")
        return []


# ----------------------------------------------------------------------------
# Cálculos
# ----------------------------------------------------------------------------
def estado_de(detalles, sucursal, camara):
    info = detalles.get(sucursal, {}).get(camara)
    return info["estado"] if info else "Por verificar"


def resumen_sucursal(detalles, sucursal):
    camaras = CATALOGO_CAMARAS[sucursal]
    op = nv = pv = 0
    for cam in camaras:
        est = estado_de(detalles, sucursal, cam)
        if est == "Operativa":
            op += 1
        elif est == "No visible":
            nv += 1
        else:
            pv += 1
    total = len(camaras)
    disp = (op / total * 100) if total else 0
    return total, op, nv, pv, disp


def resumen_total(detalles):
    total = op = nv = pv = 0
    for suc in CATALOGO_CAMARAS:
        t, o, n, p, _ = resumen_sucursal(detalles, suc)
        total, op, nv, pv = total + t, op + o, nv + n, pv + p
    disp = (op / total * 100) if total else 0
    return total, op, nv, pv, disp


def resumen_dict(detalles):
    total, op, nv, pv, disp = resumen_total(detalles)
    return {
        "total": total,
        "operativas": op,
        "no_visibles": nv,
        "por_verificar": pv,
        "disponibilidad": disp,
    }


def nivel(disp):
    if disp >= 90:
        return "Óptimo"
    if disp >= 60:
        return "Atención"
    return "Crítico"


# ----------------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------------
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_watermark_and_footer(num_pages)
            super().showPage()
        super().save()

    def draw_watermark_and_footer(self, total_pages):
        self.saveState()
        self.setFont("Helvetica-Bold", 60)
        self.setFillColor(colors.HexColor("#F2DEDE"), alpha=0.3)
        self.translate(300, 400)
        self.rotate(45)
        self.drawCentredString(0, 0, "CONFIDENCIAL")
        self.restoreState()

        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#555555"))
        self.drawString(36, 25, "INJUL - Consolidado Maestro de CCTV")
        self.drawCentredString(letter[0] / 2, 25, f"Página {self._pageNumber} de {total_pages}")
        self.drawRightString(letter[0] - 36, 25, "Documento confidencial")
        self.setStrokeColor(colors.HexColor("#CCCCCC"))
        self.setLineWidth(0.5)
        self.line(36, 38, letter[0] - 36, 38)
        self.restoreState()


COLOR_ESTADO = {
    "Operativa": "#DFF0D8",
    "No visible": "#F2DEDE",
    "Por verificar": "#FCF8E3",
}
COLOR_NIVEL = {"Óptimo": "#DFF0D8", "Atención": "#FCF8E3", "Crítico": "#F2DEDE"}


def estilo_tabla(color_cabecera, extra=None):
    comandos = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(color_cabecera)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
    ]
    if extra:
        comandos.extend(extra)
    return TableStyle(comandos)


def grafica_estado(detalles):
    """Barras horizontales apiladas: operativas / no visibles / por verificar por sucursal."""
    nombres, ops, nvs, pvs = [], [], [], []
    for suc in CATALOGO_CAMARAS:
        _, o, n, p, _ = resumen_sucursal(detalles, suc)
        nombres.append(suc)
        ops.append(o)
        nvs.append(n)
        pvs.append(p)
    maximo = max((o + n + p) for o, n, p in zip(ops, nvs, pvs))
    tope = ((maximo // 10) + 1) * 10

    dibujo = Drawing(540, 215)
    ch = HorizontalBarChart()
    ch.x, ch.y, ch.width, ch.height = 150, 22, 370, 165
    ch.data = [ops, nvs, pvs]
    ch.categoryAxis.categoryNames = nombres
    ch.categoryAxis.style = "stacked"
    ch.categoryAxis.reverseDirection = 1
    ch.categoryAxis.labels.fontName = "Helvetica"
    ch.categoryAxis.labels.fontSize = 7
    ch.categoryAxis.labels.boxAnchor = "e"
    ch.categoryAxis.labels.dx = -4
    ch.valueAxis.valueMin = 0
    ch.valueAxis.valueMax = tope
    ch.valueAxis.valueStep = 10
    ch.valueAxis.labels.fontName = "Helvetica"
    ch.valueAxis.labels.fontSize = 7
    ch.bars[0].fillColor = colors.HexColor("#2E9E2E")
    ch.bars[1].fillColor = colors.HexColor("#D64545")
    ch.bars[2].fillColor = colors.HexColor("#F0A91E")
    ch.bars.strokeColor = None
    dibujo.add(ch)

    x = 150
    for etiqueta, color in [("Operativas", "#2E9E2E"), ("No visibles", "#D64545"), ("Por verificar", "#F0A91E")]:
        dibujo.add(Rect(x, 199, 8, 8, fillColor=colors.HexColor(color), strokeColor=None))
        dibujo.add(String(x + 12, 200, etiqueta, fontName="Helvetica", fontSize=7, fillColor=colors.HexColor("#333333")))
        x += 85
    dibujo.add(String(335, 6, "Cantidad de cámaras", fontName="Helvetica", fontSize=7, textAnchor="middle", fillColor=colors.HexColor("#555555")))
    return dibujo


def hallazgos(datos):
    detalles, tickets = datos["detalles"], datos["tickets"]
    total, op, nv, pv, disp = resumen_total(detalles)
    niveles = [nivel(resumen_sucursal(detalles, s)[4]) for s in CATALOGO_CAMARAS]
    n_suc = len(CATALOGO_CAMARAS)
    lista = [
        f"Disponibilidad global de {disp:.1f}%: {op} de {total} cámaras operativas en EZStation. "
        f"{niveles.count('Crítico')} de {n_suc} sucursales están en nivel crítico y "
        f"{niveles.count('Óptimo')} en nivel óptimo."
    ]
    peor = max(CATALOGO_CAMARAS, key=lambda s: resumen_sucursal(detalles, s)[2])
    t_peor, _, nv_peor, _, _ = resumen_sucursal(detalles, peor)
    if nv_peor:
        lista.append(f"{peor}: {nv_peor} de {t_peor} cámaras no visibles en EZStation (la sucursal con más fallas).")
    genericas = sum(
        1
        for suc, camaras in CATALOGO_CAMARAS.items()
        for cam in camaras
        if estado_de(detalles, suc, cam) == "No visible" and cam.startswith("IP Camera")
    )
    if nv:
        lista.append(f"{genericas} de las {nv} cámaras no visibles conservan el nombre genérico de canal («IP Camera NN»).")
    if pv:
        donde = ", ".join(
            f"{s} ({resumen_sucursal(detalles, s)[3]})"
            for s in CATALOGO_CAMARAS
            if resumen_sucursal(detalles, s)[3]
        )
        lista.append(f"{pv} cámaras están «Por verificar»: {donde}.")
    abiertos = [t for t in tickets if t[3] == "En proceso"]
    if abiertos:
        dias = []
        for t in abiertos:
            try:
                dias.append(int(t[4]))
            except ValueError:
                pass
        extra = f"; el más antiguo lleva {max(dias)} días abierto" if dias else ""
        lista.append(f"Hay {len(abiertos)} tickets en proceso con el proveedor de {len(tickets)} registrados{extra}.")
    else:
        lista.append(f"No hay tickets en proceso ({len(tickets)} registrados).")
    return lista


def construir_pdf(datos):
    detalles = datos["detalles"]
    tickets = datos["tickets"]
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=50,
    )
    title_style = ParagraphStyle(
        "TitleStyle", fontName="Helvetica-Bold", fontSize=15, leading=19,
        textColor=colors.HexColor("#1F497D"),
    )
    seccion_style = ParagraphStyle(
        "SeccionStyle", fontName="Helvetica-Bold", fontSize=9, leading=12,
        textColor=colors.HexColor("#1F497D"),
    )
    normal_style = ParagraphStyle(
        "NormalStyle", fontName="Helvetica", fontSize=10, leading=14,
        textColor=colors.HexColor("#333333"),
    )
    celda = ParagraphStyle("Celda", fontName="Helvetica", fontSize=8, leading=10)
    story = []

    # --- PÁGINA 1: PORTADA + INDICADORES + ESTADO POR SUCURSAL + GRÁFICA ---
    fecha_txt = ahora().strftime("%d/%m/%Y")
    logo_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logo.png")
    encabezado = Paragraph(
        "<b>INFORME CONSOLIDADO DE CCTV</b><br/>"
        "<font size=8 color='#555555'>Investigadores Jurídicos EIRL - Supervisión en EZStation<br/>"
        f"Fecha: {fecha_txt} | Rango: 07:00 AM - 07:00 PM<br/>"
        "Elaborado por: Centro de Monitoreo INJUL</font>",
        title_style,
    )
    if os.path.exists(logo_path):
        logo_img = Image(logo_path, width=0.9 * inch, height=0.9 * inch)
        t_header = Table([[logo_img, encabezado]], colWidths=[75, 465])
        t_header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story.append(t_header)
    else:
        story.append(encabezado)
    story.append(Spacer(1, 10))

    total, op, nv, pv, disp = resumen_total(detalles)
    en_proceso = sum(1 for t in tickets if t[3] == "En proceso")
    pct = lambda x: f"{(x / total * 100) if total else 0:.0f}% del total"  # noqa: E731
    story.append(Paragraph("INDICADORES CLAVE", seccion_style))
    kpi = [
        ["Cámaras supervisadas", "Operativas", "No visibles", "Por verificar", "Disponibilidad", "Tickets en proceso"],
        [str(total), str(op), str(nv), str(pv), f"{disp:.1f}%", str(en_proceso)],
        [f"en {len(CATALOGO_CAMARAS)} sucursales", pct(op), pct(nv), "por confirmar", "operativas / total", f"de {len(tickets)} registrados"],
    ]
    t_kpi = Table(kpi, colWidths=[90] * 6)
    t_kpi.setStyle(
        TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 7),
            ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 1), (-1, 1), 16),
            ("FONTSIZE", (0, 2), (-1, 2), 7),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#555555")),
            ("TEXTCOLOR", (1, 1), (1, 1), colors.HexColor("#2E9E2E")),
            ("TEXTCOLOR", (2, 1), (2, 1), colors.HexColor("#D64545")),
            ("TEXTCOLOR", (3, 1), (3, 1), colors.HexColor("#C98A00")),
            ("TEXTCOLOR", (0, 1), (0, 1), colors.HexColor("#1F497D")),
            ("TEXTCOLOR", (4, 1), (5, 1), colors.HexColor("#1F497D")),
            ("TEXTCOLOR", (0, 2), (-1, 2), colors.HexColor("#777777")),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8F9FA")),
            ("LINEABOVE", (0, 0), (-1, 0), 1.5, colors.HexColor("#1F497D")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ])
    )
    story.append(t_kpi)
    story.append(Spacer(1, 10))

    story.append(Paragraph("ESTADO POR SUCURSAL", seccion_style))
    story.append(Paragraph(
        "<font size=6 color='#777777'>Disponibilidad = cámaras operativas / total de cámaras de la sucursal. "
        "Nivel: Óptimo ≥ 90 % · Atención 60-89 % · Crítico &lt; 60 %.</font>", normal_style))
    filas = [["#", "Sucursal", "Total", "Operativas", "No visibles", "Por verificar", "Disponib.", "Nivel"]]
    extra = [("ALIGN", (1, 1), (1, -1), "LEFT")]
    for i, suc in enumerate(CATALOGO_CAMARAS, start=1):
        t, o, n, p, d = resumen_sucursal(detalles, suc)
        niv = nivel(d)
        filas.append([str(i), suc, str(t), str(o), str(n), str(p), f"{d:.1f}%", niv])
        extra.append(("BACKGROUND", (7, i), (7, i), colors.HexColor(COLOR_NIVEL[niv])))
    filas.append(["", "TOTAL GENERAL", str(total), str(op), str(nv), str(pv), f"{disp:.1f}%", nivel(disp)])
    extra += [
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E9EEF5")),
        ("BACKGROUND", (7, -1), (7, -1), colors.HexColor(COLOR_NIVEL[nivel(disp)])),
    ]
    t_res = Table(filas, colWidths=[22, 150, 40, 62, 62, 70, 70, 64], repeatRows=1)
    t_res.setStyle(estilo_tabla("#1F497D", extra))
    story.append(t_res)
    story.append(Spacer(1, 10))

    story.append(Paragraph("DISTRIBUCIÓN DEL ESTADO POR SUCURSAL", seccion_style))
    story.append(grafica_estado(detalles))
    story.append(PageBreak())

    # --- PÁGINA 2: HALLAZGOS ---
    story.append(Paragraph("<b>Hallazgos del Día</b>", title_style))
    story.append(Spacer(1, 10))
    for n, texto in enumerate(hallazgos(datos), start=1):
        story.append(Paragraph(f"<b>{n}.</b> {esc(texto)}", normal_style))
        story.append(Spacer(1, 6))
    if datos.get("actualizado"):
        story.append(Spacer(1, 6))
        story.append(Paragraph(
            f"<font size=8 color='#777777'>Último cambio registrado en el sistema: {esc(datos['actualizado'])}</font>",
            normal_style))
    story.append(PageBreak())

    # --- UNA SECCIÓN POR SUCURSAL ---
    total_suc = len(CATALOGO_CAMARAS)
    for idx, (suc, camaras) in enumerate(CATALOGO_CAMARAS.items(), start=1):
        story.append(Paragraph(f"<b>Sucursal {idx} de {total_suc}: {esc(suc)}</b>", title_style))
        story.append(Spacer(1, 10))
        filas = [["#", "Ubicación / Nombre en EZStation", "En vivo", "Grabación", "Estado", "Observaciones"]]
        extra = [("ALIGN", (1, 1), (1, -1), "LEFT"), ("ALIGN", (5, 1), (5, -1), "LEFT")]
        for n, cam in enumerate(camaras, start=1):
            est = estado_de(detalles, suc, cam)
            obs = detalles.get(suc, {}).get(cam, {}).get("observacion", "")
            en_vivo = {"Operativa": "Sí", "No visible": "No"}.get(est, "Por verificar")
            grab = {"Operativa": "Normal", "No visible": "Revisar"}.get(est, "Por verificar")
            filas.append([str(n), Paragraph(esc(cam), celda), en_vivo, grab, est, Paragraph(esc(obs), celda)])
            extra.append(("BACKGROUND", (4, n), (4, n), colors.HexColor(COLOR_ESTADO.get(est, "#FCF8E3"))))
        t_suc = Table(filas, colWidths=[25, 185, 50, 70, 70, 140], repeatRows=1)
        t_suc.setStyle(estilo_tabla("#333333", extra))
        story.append(t_suc)
        story.append(PageBreak())

    # --- TICKETS ---
    story.append(Paragraph("<b>Seguimiento de Tickets con el Proveedor</b>", title_style))
    story.append(Spacer(1, 10))
    filas = [["Ticket", "Fecha", "Sucursal / Área", "Estado", "Días", "Detalle"]]
    if tickets:
        for t in tickets:
            filas.append([
                Paragraph(esc(t[0]), celda), t[1], Paragraph(esc(t[2]), celda),
                t[3], t[4], Paragraph(esc(t[5]), celda),
            ])
    else:
        filas.append(["N/A", "N/A", "Sin tickets registrados", "N/A", "0", "N/A"])
    t_tickets = Table(filas, colWidths=[60, 70, 120, 70, 40, 170], repeatRows=1)
    t_tickets.setStyle(estilo_tabla("#1F497D"))
    story.append(t_tickets)

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------
def construir_excel(datos, historial):
    detalles = datos["detalles"]
    tickets = datos["tickets"]
    azul = PatternFill("solid", fgColor="1F497D")
    blanco = Font(bold=True, color="FFFFFF")
    rellenos = {
        "Operativa": PatternFill("solid", fgColor="DFF0D8"),
        "No visible": PatternFill("solid", fgColor="F2DEDE"),
        "Por verificar": PatternFill("solid", fgColor="FCF8E3"),
    }

    def cabecera(ws, fila, titulos):
        for col, titulo in enumerate(titulos, start=1):
            c = ws.cell(row=fila, column=col, value=titulo)
            c.fill, c.font = azul, blanco
            c.alignment = Alignment(horizontal="center", vertical="center")

    wb = openpyxl.Workbook()

    # Hoja 1: resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "INJUL - Consolidado de CCTV"
    ws["A1"].font = Font(size=14, bold=True, color="1F497D")
    ws["A2"] = f"Fecha: {ahora().strftime('%d/%m/%Y')}"
    cabecera(ws, 4, ["Sucursal", "Total", "Operativas", "No visibles", "Por verificar", "Disponibilidad %", "Nivel"])
    fila = 5
    for suc in CATALOGO_CAMARAS:
        t, o, n, p, d = resumen_sucursal(detalles, suc)
        for col, valor in enumerate([suc, t, o, n, p, round(d, 1), nivel(d)], start=1):
            ws.cell(row=fila, column=col, value=valor)
        fila += 1
    t, o, n, p, d = resumen_total(detalles)
    for col, valor in enumerate(["TOTAL", t, o, n, p, round(d, 1), nivel(d)], start=1):
        ws.cell(row=fila, column=col, value=valor).font = Font(bold=True)
    for col, ancho in zip("ABCDEFG", [34, 10, 12, 13, 14, 18, 12]):
        ws.column_dimensions[col].width = ancho

    # Hoja 2: matriz de cámaras
    wm = wb.create_sheet("Matriz")
    cabecera(wm, 1, ["Sucursal", "#", "Cámara (EZStation)", "Estado", "Observación"])
    fila = 2
    for suc, camaras in CATALOGO_CAMARAS.items():
        for n, cam in enumerate(camaras, start=1):
            est = estado_de(detalles, suc, cam)
            obs = detalles.get(suc, {}).get(cam, {}).get("observacion", "")
            for col, valor in enumerate([suc, n, cam, est, obs], start=1):
                wm.cell(row=fila, column=col, value=valor)
            wm.cell(row=fila, column=4).fill = rellenos.get(est, rellenos["Por verificar"])
            fila += 1
    wm.freeze_panes = "A2"
    for col, ancho in zip("ABCDE", [30, 6, 42, 16, 50]):
        wm.column_dimensions[col].width = ancho

    # Hoja 3: tickets
    wt = wb.create_sheet("Tickets")
    cabecera(wt, 1, ["Ticket", "Fecha", "Sucursal / Área", "Estado", "Días abiertos", "Detalle"])
    for fila, t in enumerate(tickets, start=2):
        for col, valor in enumerate(t, start=1):
            wt.cell(row=fila, column=col, value=valor)
    wt.freeze_panes = "A2"
    for col, ancho in zip("ABCDEF", [14, 14, 30, 14, 14, 60]):
        wt.column_dimensions[col].width = ancho

    # Hoja 4: historial de cambios (trazabilidad)
    wh = wb.create_sheet("Historial")
    cabecera(wh, 1, ["Fecha y hora", "Sucursal", "Cámara", "Estado anterior", "Estado nuevo",
                     "Observación anterior", "Observación nueva", "Origen"])
    for fila, h in enumerate(historial, start=2):
        for col, valor in enumerate(h, start=1):
            wh.cell(row=fila, column=col, value=valor)
    wh.freeze_panes = "A2"
    for col, ancho in zip("ABCDEFGH", [24, 28, 34, 16, 16, 36, 36, 16]):
        wh.column_dimensions[col].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ----------------------------------------------------------------------------
# Rutas
# ----------------------------------------------------------------------------
@app.route("/health")
def health():
    return "ok", 200


@app.route("/")
def index():
    tipo, mensaje = MENSAJES.get(request.args.get("msg"), (None, None))
    try:
        seleccion = int(request.args.get("s", 0))
    except ValueError:
        seleccion = 0
    datos = cargar_datos()
    return render_template_string(
        HTML_TEMPLATE,
        catalogo=CATALOGO_CAMARAS,
        detalles=datos["detalles"],
        resumen=resumen_dict(datos["detalles"]),
        estados=ESTADOS,
        clases=CLASES_ESTADO,
        mensaje=mensaje,
        tipo=tipo,
        seleccion=seleccion,
    )


@app.route("/api/camara", methods=["POST"])
def api_camara():
    """Guarda al instante el estado/observación de UNA cámara y registra el cambio en el historial."""
    datos_in = request.get_json(silent=True) or {}
    sucursal = datos_in.get("sucursal")
    estado = datos_in.get("estado")
    observacion = str(datos_in.get("observacion") or "").strip()[:500]
    idx = datos_in.get("idx")

    if sucursal not in CATALOGO_CAMARAS:
        return jsonify(ok=False, error="Sucursal no válida"), 400
    if not isinstance(idx, int) or isinstance(idx, bool) or not 0 <= idx < len(CATALOGO_CAMARAS[sucursal]):
        return jsonify(ok=False, error="Cámara no válida"), 400
    if estado not in ESTADOS_VALIDOS:
        return jsonify(ok=False, error="Estado no válido"), 400
    if not DATABASE_URL:
        return jsonify(ok=False, error="Base de datos no configurada"), 503

    camara = CATALOGO_CAMARAS[sucursal][idx]
    try:
        with db_cursor() as cur:
            cur.execute(
                """
                SELECT estado, observacion FROM estado_camaras_actual
                WHERE sucursal = %s AND camara = %s FOR UPDATE
                """,
                (sucursal, camara),
            )
            previo = cur.fetchone()
            est_ant = previo[0] if previo else None
            obs_ant = (previo[1] or "") if previo else None
            cambiado = previo is None or est_ant != estado or obs_ant != observacion
            if cambiado:
                cur.execute(
                    """
                    INSERT INTO estado_camaras_actual (sucursal, camara, estado, observacion, actualizado_en)
                    VALUES (%s, %s, %s, %s, NOW())
                    ON CONFLICT (sucursal, camara) DO UPDATE
                    SET estado = EXCLUDED.estado,
                        observacion = EXCLUDED.observacion,
                        actualizado_en = NOW()
                    """,
                    (sucursal, camara, estado, observacion),
                )
                cur.execute(
                    """
                    INSERT INTO historial_cctv
                        (sucursal, camara, estado_anterior, estado_nuevo,
                         observacion_anterior, observacion_nueva, origen)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                    """,
                    (sucursal, camara, est_ant, estado, obs_ant, observacion, ip_cliente()),
                )
    except Exception:  # noqa: BLE001
        logger.exception("Error guardando cámara")
        return jsonify(ok=False, error="No se pudo guardar"), 500

    # La escritura ya está confirmada en Postgres: se invalida Redis y se recarga
    # de inmediato, así el panel y el próximo reporte ya ven el dato nuevo.
    cache_invalidar()
    datos = cargar_datos()
    return jsonify(ok=True, cambiado=cambiado, resumen=resumen_dict(datos["detalles"]))


@app.route("/guardar-ticket", methods=["POST"])
def guardar_ticket():
    ticket = (request.form.get("ticket") or "").strip()[:50]
    sucursal_area = (request.form.get("sucursal_area") or "").strip()[:150]
    estado_ticket = request.form.get("estado_ticket") or "En proceso"
    if estado_ticket not in ("En proceso", "Terminado"):
        estado_ticket = "En proceso"
    detalle = (request.form.get("detalle") or "").strip()
    try:
        dias_abiertos = max(0, int(request.form.get("dias_abiertos") or 0))
    except ValueError:
        dias_abiertos = 0

    if not ticket or not sucursal_area:
        return redirect(url_for("index", msg="error"))
    if not DATABASE_URL:
        return redirect(url_for("index", msg="sin_db"))

    try:
        with db_cursor() as cur:
            cur.execute(
                """
                INSERT INTO reportes_tickets
                    (fecha, ticket, sucursal_area, estado, dias_abiertos, detalle)
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                (ahora().strftime("%Y-%m-%d"), ticket, sucursal_area, estado_ticket, dias_abiertos, detalle),
            )
    except Exception:  # noqa: BLE001
        logger.exception("Error guardando ticket")
        return redirect(url_for("index", msg="error"))

    cache_invalidar()
    cargar_datos()
    return redirect(url_for("index", msg="ticket_ok"))


@app.route("/historial")
def historial():
    filtro = request.args.get("sucursal") or ""
    if filtro not in CATALOGO_CAMARAS:
        filtro = ""
    return render_template_string(
        HISTORIAL_TEMPLATE,
        filas=cargar_historial(200, filtro or None),
        catalogo=CATALOGO_CAMARAS,
        filtro=filtro,
    )


@app.route("/descargar-pdf")
def descargar_pdf():
    buffer = construir_pdf(cargar_datos())
    return send_file(
        buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name="Reporte_CCTV_INJUL.pdf",
    )


@app.route("/descargar-excel")
def descargar_excel():
    buffer = construir_excel(cargar_datos(), cargar_historial(5000))
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="Reporte_CCTV_INJUL.xlsx",
    )


# --- Módulo independiente: Plan de Monitoreo Cíclico y Novedades (/ciclico) ---
# Tiene sus propias tablas (ciclico_*) y rutas; sólo lee CATALOGO_CICLICO.
import ciclico  # noqa: E402

ciclico.registrar(
    app,
    catalogo=CATALOGO_CICLICO,
    db_cursor=db_cursor,
    ahora=ahora,
    local_txt=local_txt,
    ip_cliente=ip_cliente,
    database_url=DATABASE_URL,
)

# Se ejecuta al importar, para que funcione tanto con gunicorn como con python main.py
inicializar_base_de_datos()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8080))
    app.run(host="0.0.0.0", port=port)
